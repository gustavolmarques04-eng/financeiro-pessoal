"""Importa uma vez os dados da planilha ``Plano_Financeiro_Pessoal_v2.xlsx``.

O script lê a planilha, mostra tudo que encontrou e só grava depois de
confirmação. Rodando de novo, ele avisa que a importação já foi feita e não
duplica nada.

Uma decisão já está tomada e não é mais perguntada: o **saldo anterior do
Nubank (R$ 5,54) não entra na base de distribuição** — ele é o saldo inicial
do envelope Compras pessoais. Portanto::

    Recebido real          R$ 2.475,94
    Base da distribuição   R$ 2.952,21
    Saldo inicial Compras  R$     5,54

Uso::

    python import_excel.py          # interativo
    python import_excel.py --sim    # sem perguntar
    python import_excel.py --refazer
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from sqlalchemy import select

RAIZ = Path(__file__).resolve().parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from core import budget_service as budget  # noqa: E402
from core import categories as cat  # noqa: E402
from core import repositories as repo  # noqa: E402
from core.database import init_db, session_scope  # noqa: E402
from core.models import ImportLog, IncomeType, PaymentMethod  # noqa: E402
from core.period import Period  # noqa: E402
from core.utils import (  # noqa: E402
    format_brl,
    month_label,
    month_start,
    split_proportionally,
    to_cents,
)

PLANILHA_PADRAO = RAIZ.parent / "Plano_Financeiro_Pessoal_v2.xlsx"
MARCA = "Plano_Financeiro_Pessoal_v2.xlsx"

#: Descrições cujo valor é saldo de envelope, e não base de distribuição.
#: A chave é um trecho da descrição; o valor, o slug do envelope de destino.
SALDOS_DE_ENVELOPE = {"nubank": "compras"}

MAPA_TIPO = {
    "salário": IncomeType.SALARIO,
    "salario": IncomeType.SALARIO,
    "va/vr": IncomeType.VA_VR,
    "renda extra": IncomeType.RENDA_EXTRA,
    "saldo inicial": IncomeType.SALDO_INICIAL,
    "outro": IncomeType.OUTRO,
}

MAPA_MEIO = {m.value.lower(): m for m in PaymentMethod}

#: Ordem das categorias na aba Config da planilha.
SLUGS_CONFIG = (
    "independencia",
    "reserva",
    "viagem",
    "compras",
    "namorada",
    "amigos",
    "livre",
)

#: Categorias com checkbox no Dashboard da planilha (linhas 8 a 11).
SLUGS_SEPARACAO = ("independencia", "reserva", "viagem", "compras")


# --------------------------------------------------------------------------
# Leitura
# --------------------------------------------------------------------------
@dataclass
class ReceitaLida:
    """Uma linha da aba Receitas."""

    data: date
    descricao: str
    tipo: IncomeType
    valor_cents: int
    no_orcamento: bool
    obs: str | None


@dataclass
class GastoLido:
    """Uma linha da aba Gastos."""

    data: date
    descricao: str
    categoria_nome: str
    valor_cents: int
    meio: PaymentMethod
    obs: str | None


@dataclass
class Leitura:
    """Tudo que foi encontrado na planilha."""

    mes: date
    meta_reserva_cents: int
    percentuais_bp: dict[str, int]
    receitas: list[ReceitaLida] = field(default_factory=list)
    gastos: list[GastoLido] = field(default_factory=list)
    separacoes: dict[str, int] = field(default_factory=dict)
    saldos_envelope: dict[str, int] = field(default_factory=dict)
    reserva_cents: int = 0
    investimentos_cents: int = 0
    dividendos_cents: int = 0

    @property
    def recebido_cents(self) -> int:
        """Soma das receitas que não são saldo inicial."""
        return sum(
            r.valor_cents for r in self.receitas if r.tipo is not IncomeType.SALDO_INICIAL
        )

    @property
    def base_cents(self) -> int:
        """Soma das receitas que entram na base de distribuição."""
        return sum(r.valor_cents for r in self.receitas if r.no_orcamento)


def _cents(valor: object) -> int:
    """Converte uma célula numérica em centavos, tolerando vazio."""
    if valor in (None, ""):
        return 0
    return to_cents(Decimal(str(valor)))


def _data(valor: object, padrao: date) -> date:
    """Converte uma célula de data, caindo no padrão quando vazia."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return padrao


def _envelope_de(descricao: str) -> str | None:
    """Slug do envelope quando a descrição indica saldo de envelope."""
    texto = descricao.lower()
    for pista, slug in SALDOS_DE_ENVELOPE.items():
        if pista in texto:
            return slug
    return None


def ler_planilha(caminho: Path) -> Leitura:
    """Lê a planilha simplificada de 4 abas e devolve o que encontrou.

    Aplica de uma vez a interpretação correta dos saldos de envelope: eles
    ficam fora do rateio e viram saldo inicial da categoria correspondente.
    """
    wb = openpyxl.load_workbook(caminho, data_only=True)
    dash, cfg = wb["Dashboard"], wb["Config"]

    mes = month_start(_data(dash["A4"].value, date.today()))

    percentuais = {}
    for indice, slug in enumerate(SLUGS_CONFIG):
        bruto = cfg.cell(6 + indice, 2).value or 0
        percentuais[slug] = int(round(float(bruto) * 10_000))

    leitura = Leitura(
        mes=mes,
        meta_reserva_cents=_cents(cfg["B3"].value),
        percentuais_bp=percentuais,
        reserva_cents=_cents(dash["B20"].value),
        investimentos_cents=_cents(dash["B21"].value),
    )

    receitas = wb["Receitas"]
    for linha in range(2, receitas.max_row + 1):
        descricao = receitas.cell(linha, 2).value
        valor = receitas.cell(linha, 4).value
        if not descricao or valor in (None, ""):
            continue

        texto = str(descricao).strip()
        tipo_txt = str(receitas.cell(linha, 3).value or "Outro").strip().lower()
        valor_cents = _cents(valor)
        envelope = _envelope_de(texto)

        if envelope is not None:
            # Não é renda nem base: é saldo que já existia no envelope.
            leitura.saldos_envelope[envelope] = (
                leitura.saldos_envelope.get(envelope, 0) + valor_cents
            )
            no_orcamento = False
        else:
            orcamento_txt = str(receitas.cell(linha, 5).value or "Sim").strip().lower()
            no_orcamento = not orcamento_txt.startswith("n")

        leitura.receitas.append(
            ReceitaLida(
                data=_data(receitas.cell(linha, 1).value, mes),
                descricao=texto,
                tipo=MAPA_TIPO.get(tipo_txt, IncomeType.OUTRO),
                valor_cents=valor_cents,
                no_orcamento=no_orcamento,
                obs=(
                    str(receitas.cell(linha, 6).value).strip()
                    if receitas.cell(linha, 6).value
                    else None
                ),
            )
        )

    gastos = wb["Gastos"]
    for linha in range(2, gastos.max_row + 1):
        descricao = gastos.cell(linha, 2).value
        valor = gastos.cell(linha, 4).value
        if not descricao or valor in (None, ""):
            continue
        meio_txt = str(gastos.cell(linha, 5).value or "Outro").strip().lower()
        leitura.gastos.append(
            GastoLido(
                data=_data(gastos.cell(linha, 1).value, mes),
                descricao=str(descricao).strip(),
                categoria_nome=str(gastos.cell(linha, 3).value or "Outro").strip(),
                valor_cents=_cents(valor),
                meio=MAPA_MEIO.get(meio_txt, PaymentMethod.OUTRO),
                obs=(
                    str(gastos.cell(linha, 6).value).strip()
                    if gastos.cell(linha, 6).value
                    else None
                ),
            )
        )

    # Separações confirmadas no dashboard (linhas 8 a 11, coluna "Feito?").
    # Os valores são recalculados a partir da base correta, para ficarem
    # coerentes com o rateio que o aplicativo vai mostrar.
    marcadas = [
        slug
        for indice, slug in enumerate(SLUGS_SEPARACAO)
        if "feito" in str(dash.cell(8 + indice, 4).value or "").lower()
    ]
    if marcadas:
        pesos = [percentuais[s] for s in SLUGS_CONFIG]
        fatias = dict(zip(SLUGS_CONFIG, split_proportionally(leitura.base_cents, pesos)))
        leitura.separacoes = {slug: fatias[slug] for slug in marcadas}

    return leitura


# --------------------------------------------------------------------------
# Apresentação
# --------------------------------------------------------------------------
def mostrar(leitura: Leitura) -> None:
    """Imprime tudo que foi lido, já com a interpretação aplicada."""
    print("=" * 68)
    print(f"PLANILHA LIDA — mês {month_label(leitura.mes)}")
    print("=" * 68)

    print("\nRECEITAS")
    for r in leitura.receitas:
        if _envelope_de(r.descricao):
            marca = "saldo de envelope (fora do rateio)"
        else:
            marca = "no rateio" if r.no_orcamento else "fora do rateio"
        print(
            f"  {r.data:%d/%m/%Y}  {r.descricao:<26} {r.tipo.value:<14} "
            f"{format_brl(r.valor_cents):>13}  ({marca})"
        )
    print(f"  {'Recebido no mês':<40} {format_brl(leitura.recebido_cents):>13}")
    print(f"  {'Base de distribuição':<40} {format_brl(leitura.base_cents):>13}")

    print("\nSALDOS INICIAIS DE ENVELOPE")
    if not leitura.saldos_envelope:
        print("  (nenhum)")
    for slug, valor in leitura.saldos_envelope.items():
        print(f"  {slug:<28} {format_brl(valor):>13}")

    print("\nGASTOS")
    if not leitura.gastos:
        print("  (nenhum)")
    for g in leitura.gastos:
        print(
            f"  {g.data:%d/%m/%Y}  {g.descricao:<26} {g.categoria_nome:<20} "
            f"{format_brl(g.valor_cents):>13}"
        )

    print("\nSEPARAÇÕES MARCADAS COMO FEITAS")
    if not leitura.separacoes:
        print("  (nenhuma)")
    for slug, valor in leitura.separacoes.items():
        print(f"  {slug:<28} {format_brl(valor):>13}")

    print("\nFECHAMENTO")
    print(f"  {'Reserva de emergência':<28} {format_brl(leitura.reserva_cents):>13}")
    print(f"  {'Investimentos':<28} {format_brl(leitura.investimentos_cents):>13}")
    print(f"  {'Dividendos':<28} {format_brl(leitura.dividendos_cents):>13}")

    print("\nCONFIGURAÇÃO")
    print(f"  {'Meta da reserva':<28} {format_brl(leitura.meta_reserva_cents):>13}")
    total_bp = sum(leitura.percentuais_bp.values())
    for slug, bp in leitura.percentuais_bp.items():
        print(f"  {slug:<28} {bp / 100:>12.2f}%")
    print(f"  {'TOTAL':<28} {total_bp / 100:>12.2f}%")


# --------------------------------------------------------------------------
# Gravação
# --------------------------------------------------------------------------
def ja_importado() -> bool:
    """Se a planilha já foi importada alguma vez."""
    with session_scope() as session:
        return (
            session.scalar(select(ImportLog).where(ImportLog.source == MARCA)) is not None
        )


def limpar_marca() -> None:
    """Remove a marca de importação (para reimportar de propósito)."""
    with session_scope() as session:
        registro = session.scalar(select(ImportLog).where(ImportLog.source == MARCA))
        if registro:
            session.delete(registro)


def importar(leitura: Leitura) -> None:
    """Grava tudo em uma única transação."""
    with session_scope() as session:
        ids = {v.slug: v.id for v in cat.resolve_all(session, leitura.mes)}

        repo.upsert_settings_version(
            session,
            effective_month=leitura.mes,
            meta_reserva_cents=leitura.meta_reserva_cents,
            percentuais_bp={
                ids[slug]: bp for slug, bp in leitura.percentuais_bp.items() if slug in ids
            },
        )

        for r in leitura.receitas:
            repo.create_income(
                session,
                on=r.data,
                description=r.descricao,
                type_=r.tipo,
                amount_cents=r.valor_cents,
                counts_in_budget=r.no_orcamento,
                note=r.obs,
            )

        nomes = {v.name.lower(): v.id for v in cat.resolve_all(session, leitura.mes)}
        for g in leitura.gastos:
            repo.create_expense(
                session,
                purchase_date=g.data,
                description=g.descricao,
                category_id=nomes.get(g.categoria_nome.lower(), ids["outro"]),
                total_cents=g.valor_cents,
                payment_method=g.meio,
                note=g.obs,
            )

        revisao = repo.get_revision(session, leitura.mes)
        for slug, valor in leitura.separacoes.items():
            if slug in ids:
                repo.set_allocation(
                    session, leitura.mes, ids[slug], separated_cents=valor, revision=revisao
                )

        if any(
            (leitura.reserva_cents, leitura.investimentos_cents, leitura.dividendos_cents)
        ):
            repo.upsert_closing(
                session,
                leitura.mes,
                reserva_cents=leitura.reserva_cents,
                investimentos_cents=leitura.investimentos_cents,
                note="Importado da planilha",
            )

        for slug, valor in leitura.saldos_envelope.items():
            if slug in ids:
                repo.set_opening_balance(
                    session, ids[slug], valor, note="Saldo trazido da planilha"
                )

        session.add(
            ImportLog(
                source=MARCA,
                detail=(
                    f"mês {month_label(leitura.mes)}; "
                    f"{len(leitura.receitas)} receitas; {len(leitura.gastos)} gastos; "
                    f"base {format_brl(leitura.base_cents)}"
                ),
            )
        )


def _preparar_console() -> None:
    """Garante que acentos e símbolos apareçam mesmo em console cp1252."""
    for fluxo in (sys.stdout, sys.stderr):
        reconfigurar = getattr(fluxo, "reconfigure", None)
        if reconfigurar is not None:
            reconfigurar(encoding="utf-8", errors="replace")


def main() -> int:
    """Ponto de entrada do script."""
    _preparar_console()
    parser = argparse.ArgumentParser(description="Importa a planilha para o banco.")
    parser.add_argument("--arquivo", type=Path, default=PLANILHA_PADRAO)
    parser.add_argument("--sim", action="store_true", help="não pedir confirmação")
    parser.add_argument("--refazer", action="store_true", help="permite importar de novo")
    args = parser.parse_args()

    if not args.arquivo.exists():
        print(f"Planilha não encontrada: {args.arquivo}")
        return 1

    init_db()

    if args.refazer:
        limpar_marca()
    if ja_importado():
        print("Esta planilha já foi importada. Nada foi alterado.")
        print("Use --refazer se quiser importar de novo (pode duplicar dados).")
        return 0

    leitura = ler_planilha(args.arquivo)
    mostrar(leitura)

    if not args.sim:
        confirmacao = input("\nGravar estes dados no banco? [s/N]: ").strip().lower()
        if not confirmacao.startswith("s"):
            print("Importação cancelada. Nada foi gravado.")
            return 0

    importar(leitura)
    print("\n✔ Importação concluída.")

    with session_scope() as session:
        plano = budget.get_month_plan(session, leitura.mes)
        print(f"\nResumo de {month_label(plano.month)}:")
        print(f"  Recebido no mês        : {format_brl(plano.recebido_cents)}")
        print(f"  Base de distribuição   : {format_brl(plano.base_cents)}")
        for linha in plano.separacoes:
            print(
                f"  {linha.categoria.name:<28} planejado "
                f"{format_brl(linha.planejado_cents):>13} · separado "
                f"{format_brl(linha.separado_cents):>13} · {linha.status}"
            )
        for envelope in budget.envelopes(session, Period.of_month(plano.month)):
            print(
                f"  envelope {envelope.categoria.name:<19} saldo "
                f"{format_brl(envelope.saldo_cents):>13}"
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
